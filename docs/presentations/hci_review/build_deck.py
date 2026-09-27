"""Build the HCI v1 review deck (python-pptx).

Style follows the owner's template: black title bar (Arial 26 bold, white),
black bottom line (Arial 20, white), all other text Arial >= 18 pt, and the
graphic occupying the maximum area. SVG diagrams from the code documentation
are embedded as native SVG with a PNG fallback.

Run from the repository root after ``build_assets.py`` and ``capture_ui.mjs``::

    python docs/presentations/hci_review/build_deck.py
"""

from __future__ import annotations

import json
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
A = HERE / "assets"
OUT = HERE / "hci_v1_review.pptx"

W, H = 13.333, 7.5
TITLE_H, BOTTOM_H = 0.78, 0.72
BODY_TOP, BODY_BOTTOM = TITLE_H + 0.12, H - BOTTOM_H - 0.1
FONT = "Arial"
BLACK, WHITE = RGBColor(0, 0, 0), RGBColor(255, 255, 255)
INK = RGBColor(0x1F, 0x2A, 0x33)
MAROON, TEAL, OCHRE, NAVY, RED = RGBColor(0x8B, 0x1E, 0x3F), RGBColor(0x0E, 0x7C, 0x86), RGBColor(0xB8, 0x6E, 0x00), RGBColor(0x1D, 0x35, 0x57), RGBColor(0xC1, 0x12, 0x1F)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
_svg_counter = [0]


def ppt_safe_svg(src: Path) -> Path:
    """Copy of an SVG without filter effects: PowerPoint's SVG renderer drops filtered groups."""

    import re

    dst = A / f"{src.stem}_ppt.svg"
    text = src.read_text(encoding="utf-8")
    dst.write_text(re.sub(r'\s+filter="url\(#[^)]*\)"', "", text), encoding="utf-8")
    return dst


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------


def _bar(slide, top: float, height: float, text: str, size: int, bold: bool) -> None:
    shp = slide.shapes.add_shape(1, Inches(0), Inches(top), Inches(W), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = BLACK
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, WHITE


def new_slide(title: str, bottom: str, notes: str = ""):
    slide = prs.slides.add_slide(BLANK)
    _bar(slide, 0, TITLE_H, title, 26, True)
    _bar(slide, H - BOTTOM_H, BOTTOM_H, bottom, 20, False)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def picture(slide, path: Path, left: float, top: float, width: float, height: float, svg: Path | None = None):
    """Place an image fitted and centred in the box; optionally attach an SVG (PNG is the fallback)."""

    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(width / iw, height / ih)
    w, h = iw * scale, ih * scale
    x, y = left + (width - w) / 2, top + (height - h) / 2
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if svg is not None:
        _attach_svg(slide, pic, svg)
    return pic


def _attach_svg(slide, pic, svg: Path) -> None:
    _svg_counter[0] += 1
    part = Part(PackURI(f"/ppt/media/hci_diagram{_svg_counter[0]}.svg"), "image/svg+xml", slide.part.package, svg.read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.xpath(".//a:blip")[0]
    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ext_lst = etree.SubElement(blip, f"{{{ns_a}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{ns_a}}}ext", uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip", nsmap={"asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"})
    svg_blip.set(f"{{{ns_r}}}embed", rid)


def textbox(slide, left: float, top: float, width: float, height: float, blocks: list[tuple[str, list[str], RGBColor]], size: int = 18, gap: int = 8):
    """Section headings (bold, coloured) each followed by bullet lines."""

    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    first = True
    for heading, bullets, color in blocks:
        if heading:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(0 if p is tf.paragraphs[0] else 10)
            r = p.add_run()
            r.text = heading
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size + 2), True, color
        for b in bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(gap / 2)
            r = p.add_run()
            r.text = "• " + b
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), INK
    return tb


def table(slide, left: float, top: float, width: float, rows: list[list[str]], col_w: list[float], size: int = 18, header_color: RGBColor = NAVY, row_h: float = 0.42, bold_col0: bool = False):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_color if i == 0 else (RGBColor(0xF2, 0xF5, 0xF8) if i % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = value
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = i == 0 or (bold_col0 and j == 0)
            r.font.color.rgb = WHITE if i == 0 else INK
    return tbl


def full(slide, path: Path, svg: Path | None = None, margin: float = 0.25):
    return picture(slide, path, margin, BODY_TOP, W - 2 * margin, BODY_BOTTOM - BODY_TOP, svg)


# ---------------------------------------------------------------------------
# Data for tables
# ---------------------------------------------------------------------------

cand = json.loads((ROOT / "docs/hci_candidate_study.summary.json").read_text(encoding="utf-8"))
fixed = json.loads((ROOT / "docs/hci_evidence/delta_max_rules/summary_fixed10.json").read_text(encoding="utf-8"))
auto = json.loads((ROOT / "docs/hci_evidence/summary.json").read_text(encoding="utf-8"))
real = json.loads((ROOT / "docs/hci_evidence/real_report_images/report.json").read_text(encoding="utf-8"))


def order_cell(s: dict, fam: str, key: str) -> str:
    a = s["ordering_axioms"][fam][key]
    return f"{a['replicates_monotone']}/{a['replicates']}"


def max_change(s: dict, kinds: list[str], key: str) -> str:
    inv = s["invariance_relative_change"]
    vals = [abs(inv[b][k][key]) for b in inv for k in kinds if inv[b][k].get(key) is not None]
    return f"{100 * max(vals):.1f} %" if vals else "n/a"


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

# 1 Title
s = new_slide(
    "Hydride Connectivity Index (HCI) v1: problem, method, validation, results",
    "A bounded, threshold-free connectivity index validated at constant area fraction (MicroSeg 2.0.0)",
    "Review deck for the HCI v1 formulation. It covers the problem, the intern prototype and why it was not adopted, the synthetic benchmark and its rationale, the algorithm step by step, validation evidence, the owner's delta_max rule and its consequences, results on the report's real micrographs, the implementation and the caveats.",
)
picture(s, A / "hci_concept.png", 0.35, BODY_TOP + 0.05, W - 0.7, BODY_BOTTOM - BODY_TOP - 0.55, ppt_safe_svg(A / "hci_concept.svg"))
tb = s.shapes.add_textbox(Inches(0.35), Inches(BODY_BOTTOM - 0.48), Inches(W - 0.7), Inches(0.45))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "MicroSeg 2.0.0  ·  formulation hci.v1  ·  review deck, 27 September 2026"
r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(18), RGBColor(0x55, 0x66, 0x77)

# 2 Problem
s = new_slide(
    "The problem: amount and orientation do not describe connectivity",
    "Same area fraction, same Fn — very different connectivity: a new descriptor is needed",
    "All three synthetic masks have area fraction 5.0 % and length-weighted Fn = 1.00, because every hydride is radial. Yet a crack path through the continuous lines crosses far less matrix than one through isolated platelets. Area fraction and Fn cannot distinguish them; the HCI radial (fixed delta_max = 10 um) does: 0.20, 0.44, 0.55.",
)
full(s, A / "problem_same_af.png")

# 3 Prior descriptors
s = new_slide(
    "Existing hydride descriptors and where the HCI fits",
    "None separates connectivity from amount under controlled tests — that is the gap v1 fills",
    "RHCF: Billone, Burtseva, Einziger, J. Nucl. Mater. 433 (2013) 431. HCC: Kim et al., J. Nucl. Mater. 456 (2015) 235. RHCP: Simon et al., J. Nucl. Mater. 547 (2021) 152817; Dijkstra implementation PROPHET, J. Nucl. Mater. 567 (2022). Our companion path continuity Pi is explicitly RHCP-type; novelty is claimed for the integrated linkage HCI, the critical linking distance, the closure identity and the constant-Af benchmark with axioms.",
)
rows = [
    ["Descriptor", "What it measures", "Limitation for connectivity"],
    ["RHF / Fn", "Share of radial hydrides", "Orientation only; no linkage"],
    ["HCC (Kim 2015)", "Radial length of near-continuous hydrides in a band", "Empirical band and closeness rules"],
    ["RHCF (Billone 2013)", "Max radial hydride length in an arc window", "Window-based; one direction"],
    ["RHCP (Simon 2021)", "Optimal through-wall crack path (toughness-weighted)", "Single weakest path; toughness ratio assumed"],
    ["Intern HCI prototype", "Topology × length / nearest-cluster distance", "Unbounded, amount-driven, fails invariance"],
    ["HCI v1 (this work)", "Integrated single-linkage coverage, δ½, closure κ", "Validated on synthetic axioms; field data pending"],
]
table(s, 0.45, BODY_TOP + 0.35, W - 0.9, rows, [2.6, 5.3, 4.53], row_h=0.66, bold_col0=True)

# 4 Prototype and confound
s = new_slide(
    "The intern prototype and what its results really show",
    "Prototype HCI rose with hydrogen — but so did the hydride amount: the trend is confounded",
    "Prototype: w_i = ((B_i+1)/N_i)(Lbar/d_min,i), HCI = sum(A_i w_i)/sum(A_i). Reported values 2.59, 3.38, 8.02 for 25, 70, 200 ppm. The report's own masks have area fractions 5.3, 9.1 and 13.7 %, so the rise cannot be attributed to connectivity. Re-running the unchanged prototype on the report masks gives 2.44, 3.12, 6.97 (figure resolution).",
)
textbox(s, 0.35, BODY_TOP + 0.1, 3.9, 5.6, [
    ("Prototype formula", ["w = (B+1)/N × L̄ / d_min", "HCI = Σ A·w / Σ A", "B junctions, N = ends + B", "d_min: nearest cluster"], MAROON),
    ("Reported (Table 2)", ["25 ppm: 2.59", "70 ppm: 3.38", "200 ppm: 8.02"], TEAL),
    ("But", ["Area fraction 5 → 14 %", "Unbounded, px units"], OCHRE),
])
picture(s, A / "prototype_confound.png", 4.3, BODY_TOP, W - 4.55, BODY_BOTTOM - BODY_TOP)

# 5 Prototype failures
s = new_slide(
    "The prototype on the constant-area-fraction benchmark",
    "Formulation-level failures, not just coding bugs — the prototype was not adopted",
    "All masks at Af = 5 %. (a) HCI falls when collinear chains finally merge. (b) A single connected cluster has no neighbour distance, so rings and the fully connected honeycomb mesh score 0. (c) Identical lines rotated vary by 15 % CV although the formula is isotropic. (d) Transpose changes it by 7.7 %, resolution by up to 6.8 %, roughness by 19 %. Also: two code paths disagree (6.05 vs 4.41), and six percolating masks exceeded 300 s.",
)
full(s, A / "prototype_failures.png")

# 6 Axioms
s = new_slide(
    "Ten testable design axioms for any connectivity index",
    "Every axiom is checked automatically on the synthetic benchmark (tests/test_phase38_hci_core.py)",
    "The axioms turn the scientific intent into pass/fail tests. A1 is the key one: at constant area fraction the index must rank arrangements by connectivity.",
)
rows = [
    ["ID", "Requirement", "ID", "Requirement"],
    ["A1", "Ranks connectivity at constant Af", "A6", "Exact symmetries (flip, transpose)"],
    ["A2", "Bounded in (0, 1], interpretable", "A7", "Resolution-consistent (µm thresholds)"],
    ["A3", "Linking never lowers it", "A8", "Robust to segmentation noise (≤ 5 %)"],
    ["A4", "Fragmentation never raises it", "A9", "Directional where declared"],
    ["A5", "Defined for empty and single clusters", "A10", "Deterministic, full provenance"],
]
table(s, 0.6, BODY_TOP + 0.6, W - 1.2, rows, [0.9, 5.17, 0.9, 5.16], row_h=0.8)

# 7 Benchmark construction
s = new_slide(
    "Synthetic benchmark: construction and rationale",
    "Controlled masks with exact known answers make code and concept validation objective",
    "Each sample is a centreline graph (nodes, straight edges) rasterised as 4 px capsules. The graph gives exact ground truth: length, ends, junction degrees, loops, components and closure. A size parameter is solved so that the raster area fraction is 5.00 ± 0.05 %. Separate objects keep at least 12 px (6 um) of matrix. 203 masks regenerate bit-for-bit (tested).",
)
textbox(s, 0.35, BODY_TOP + 0.1, 4.4, 5.6, [
    ("Why synthetic?", ["Real specimens change amount, orientation and spacing together", "One factor varied at a time"], MAROON),
    ("How", ["Centreline graph → 4 px capsules", "Af fixed at 5.00 ± 0.05 %", "≥ 6 µm clearance between objects"], TEAL),
    ("Ground truth", ["Ends, junction degrees, loops", "203 masks, bit-for-bit reproducible"], OCHRE),
])
picture(s, A / "benchmark_construction.png", 4.8, BODY_TOP, W - 5.05, BODY_BOTTOM - BODY_TOP)


def family_slide(title: str, bottom: str, notes: str, items: list[tuple[str, str, str]]):
    s = new_slide(title, bottom, notes)
    n = len(items)
    avail = BODY_BOTTOM - BODY_TOP
    slot = avail / n
    for i, (asset, head, why) in enumerate(items):
        top = BODY_TOP + i * slot
        tb = s.shapes.add_textbox(Inches(0.3), Inches(top + 0.05), Inches(3.3), Inches(slot - 0.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = head
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(20), True, MAROON
        p = tf.add_paragraph()
        r = p.add_run(); r.text = why
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(18), INK
        picture(s, A / asset, 3.6, top + 0.02, W - 3.85, slot - 0.08)
    return s


family_slide("Benchmark families 1: gap and fragmentation (Af = 5 %)", "Expected: closing gaps raises connectivity; splitting the same hydride lowers it",
             "Gap sweep: 10 chains of 3 collinear radial segments; surface gap 40, 20, 10, 6, 3 px, then merged. Isolates linking of collinear hydrides. Fragmentation: the same total length split into 8 to 128 radial pieces. Isolates fragment length at fixed amount.",
             [("family_gap_sweep.png", "Gap sweep", "Only the gap between collinear hydrides changes"), ("family_fragmentation.png", "Fragmentation", "Same hydride split into more, shorter pieces")])
family_slide("Benchmark families 2: arrangement and orientation", "Expected: aligned < stringers < stepped ≈ continuous; radial index follows angle, isotropic stays flat",
             "Arrangement: 36 equal segments placed randomly (circumferential, isotropic, radial), as stringers with 3 um gaps, as stepped linked staircases, or merged into continuous lines. Orientation: 12 identical lines rotated 0 to 90 degrees from the circumferential axis.",
             [("family_arrangement.png", "Arrangement", "Same segments, different spatial organisation"), ("family_orientation.png", "Orientation", "Identical lines, rotated 0–90°")])
family_slide("Benchmark families 3: topology and the percolation (Boolean) model", "Expected: topology separates ends from loops; HCI rises with Af across the percolation transition",
             "Topology: equal total length as lines, Y (degree 3), X (degree 4), branched trees, hexagonal rings, and one honeycomb mesh; exact ends, junction degrees and loops known. Boolean model: overlapping random 60 px sticks at Af = 2 to 20 %, isotropic and radial, the classical continuum-percolation system; this family deliberately varies Af.",
             [("family_topology.png", "Topology", "Ends, junctions, loops and one mesh"), ("family_boolean_model.png", "Boolean model", "Af varied 2–20 %: percolation check")])
family_slide("Benchmark families 4: invariance, noise and edge cases", "Expected: exact symmetry, ≤ 5 % change under resolution or noise, explicit status for edge cases",
             "Invariance: transpose, flips, resolution x0.5 and x2 (pixel size recorded), boundary roughness (10 % boundary pixel flips), speckle below the clean-up size, pinholes, spurs and 2-px breaks. Degenerate: empty, sub-threshold speck, a single segment, a border-to-border line, a ring, a single mesh cluster and a pair of lines.",
             [("family_invariance.png", "Invariance", "Symmetry, resolution and segmentation noise"), ("family_degenerate.png", "Edge cases", "Empty, speck, single clusters")])

# How it works
s = new_slide("How the HCI works", "Link hydrides across matrix gaps ≤ δ, measure covered length, average over δ up to δ_max",
              "Components are linked when the matrix ligament between them is at most delta (single linkage; bridges are synthetic and add no area). C(delta) is the area-weighted fraction of the reference length covered by the cluster each hydride belongs to, using projected coverage so gaps add no extent. The HCI is the mean of C over 0 to delta_max: bounded, monotone and threshold-free inside the range.")
full(s, A / "hci_concept.png", ppt_safe_svg(A / "hci_concept.svg"))
s = new_slide("Algorithm flow sheet (same SVG as the code documentation)", "One deterministic core for CLI, web and Python; real hydride and synthetic bridges never mixed",
              "Steps: binarise and clean (5 um2, 2x2 px floor); label 8-connected components; gap graph from the Euclidean distance transform (Voronoi neighbours contain the MST, so single linkage is exact); Kruskal merges give the exact step function C_u(delta); exact integral; companions path continuity and skeleton topology; quality flags; result schema microseg.hci_result.v1.")
full(s, A / "hci_v1_algorithm_flow.png", ppt_safe_svg(A / "hci_v1_algorithm_flow.svg"))

steps = [
    ("syn_stage1.png", "Step 1–2: clean-up and components", "The mask is cleaned and split into 8-connected hydride components (36 platelets)", "Input: arrangement__radial_stringers__r00 (Af 5 %, 0.5 um/px). Components below 5 um2 would be removed; holes below 5 um2 filled; a 2x2 px floor applies at coarse resolution."),
    ("syn_stage2.png", "Step 3: gap graph from the distance transform", "Only Voronoi-neighbour gaps are needed — they contain the minimum spanning tree (exact, O(pixels))", "Background colours show the nearest component of every matrix pixel. Red lines are minimum-spanning gap edges with their surface gap in um; the 3 um stringer gaps are the shortest."),
    ("syn_stage3.png", "Step 4: clusters as the association distance δ grows", "Stringer gaps (3 µm) bridge first: C_R jumps from 0.18 to 0.54 — red lines are synthetic bridges", "Clusters are single-linkage components at each delta. Bridges add no area and are drawn in red only in overlays."),
    ("syn_stage4.png", "Step 5: coverage, connectivity function and the HCI", "HCI = area under C(δ) ÷ δ_max = 0.44; δ½ = 3.2 µm recovers the designed gap", "Left: the largest cluster and the rows its hydride covers (bridged gaps add none). Right: exact step functions for three directions; the shaded area divided by delta_max is HCI radial; dots mark delta_half."),
    ("syn_stage5.png", "Companions: weakest path and skeleton topology", "Π (RHCP-type path) and κ (network closure) are reported beside the HCI, never multiplied in", "Path continuity: two Dijkstra sweeps; ligament of the cheapest edge-to-edge path through each pixel. Topology: ends, collapsed junctions and loops from the Euler number; identity N_E = beta + 2C - 2mu holds exactly (kappa = 0 for these free-ended stringers)."),
]
for asset, title, bottom, notes in steps:
    s = new_slide(title, bottom, notes)
    full(s, A / asset)

# Evidence
s = new_slide("Validation on the benchmark: HCI v1 (fixed δ_max = 10 µm)", "v1 meets every ordering axiom in every replicate at constant area fraction",
              "Mean +/- SD over replicates. HCI radial rises strictly as gaps close (0.22 to 0.66), falls with fragmentation, separates random (0.20), stringers (0.44), stepped (0.53) and continuous (0.55), follows orientation, and gives the mesh the maximum. The Boolean model shows the expected percolation rise with Af.")
full(s, A / "benchmark_evidence_fixed10.png")

s = new_slide("Scorecard: prototype versus HCI v1", "v1 fixes ordering, symmetry, resolution and single-cluster failures; noise sensitivity is similar",
              "Prototype from the unchanged intern script (candidate study). v1 from the production library with fixed delta_max = 10 um. Ordering = replicates in the expected order (tolerance 0.01). Changes are maximum relative changes over the three invariance reference samples.")
rows = [
    ["Test", "Prototype", "HCI v1 (fixed 10 µm)"],
    ["Gap sweep ordering", order_cell(cand, "gap_sweep", "proto_workbook"), order_cell(fixed, "gap_sweep", "HCI_R")],
    ["Fragmentation ordering", order_cell(cand, "fragmentation", "proto_workbook"), order_cell(fixed, "fragmentation", "HCI_R")],
    ["Arrangement ordering", order_cell(cand, "arrangement", "proto_workbook"), order_cell(fixed, "arrangement", "HCI_R")],
    ["Orientation ordering", order_cell(cand, "orientation", "proto_workbook") + " (should be flat)", order_cell(fixed, "orientation", "HCI_R")],
    ["Transpose / flips", max_change(cand, ["transpose", "flip_lr", "flip_ud"], "proto_workbook"), max_change(fixed, ["transpose", "flip_lr", "flip_ud"], "HCI_R")],
    ["Resolution ×0.5, ×2", max_change(cand, ["scale_0.5", "scale_2.0"], "proto_workbook"), max_change(fixed, ["scale_0.5", "scale_2.0"], "HCI_R")],
    ["Speckle, pinholes, spurs, breaks", max_change(cand, ["speckle", "pinholes", "spurs", "breaks"], "proto_workbook"), max_change(fixed, ["speckle", "pinholes", "spurs", "breaks"], "HCI_R")],
    ["Boundary roughness (10 % flips)", max_change(cand, ["boundary_roughness"], "proto_workbook"), max_change(fixed, ["boundary_roughness"], "HCI_R")],
    ["Single cluster / mesh", "0 (undefined)", "finite (iso 1.00)"],
]
table(s, 1.2, BODY_TOP + 0.1, W - 2.4, rows, [4.7, 3.7, 2.53], row_h=0.52, bold_col0=True)

s = new_slide("The δ_max rule: owner default versus alternatives", "Default kept (1/5 of smallest hydride); compare specimens only with one fixed δ_max",
              "Owner decision: delta_max = 1/5 of the smallest accepted hydride. It is scale-free, but it is recomputed per image, so different images are integrated over different ranges (arrangement 3/5, orientation 1/3), and one small fragment sets the range for the whole image (-61 % under 2-px breaks). A median rule fixes the break sensitivity (2.4 %) but not comparability; fixed 10 um passes everything. The web UI offers all three; the Help and docs require a fixed value for comparisons.")
full(s, A / "delta_rules.png")

# Real images
s = new_slide("Real micrographs from the intern report (25, 70, 200 ppm)", "The same pipeline runs on the report masks and on re-segmented micrographs (ML model)",
              "Images are the report's Table 2 figures: optical 513x384 (100 um bar = 46.5 px, 2.15 um/px) and the intern's masks 344x258 (3.21 um/px). Columns: optical micrograph; intern mask; HCI clusters at delta_max = 10 um (red = bridges); MicroSeg ML re-segmentation of the optical image with its clusters. Figure resolution only, so values are indicative.")
full(s, A / "real_grid.png")
for asset, title, bottom, notes in [
    ("real200_stage2.png", "200 ppm, step 3: gap graph on a real mask", "On real masks the gap graph is dense; each gap sets the δ at which two clusters merge", "Nearest-component map and minimum-spanning gap edges (um) for the 200 ppm report mask at 3.21 um/px."),
    ("real200_stage3.png", "200 ppm, step 4: clusters as δ grows", "Circumferential chains link first; radial linking needs much larger gaps", "Clusters at delta = 0, 5, 10 and 20 um; the number of clusters falls 130 → 20 while C_R stays low because the clusters run circumferentially."),
    ("real200_stage4.png", "200 ppm, step 5: coverage and connectivity functions", "Circumferential and isotropic connectivity are high; radial connectivity stays low (HCI R = 0.06)", "Largest cluster coverage and the three connectivity step functions for the 200 ppm mask; delta_max = 10 um."),
    ("real200_stage5.png", "200 ppm, companions: weakest path and topology", "The weakest radial path must cross mostly matrix; the network starts to close (κ = 0.13)", "Pi_R and skeleton topology for the 200 ppm mask; identity residual 0 after the resolution floor and bookkeeping fixes."),
]:
    s = new_slide(title, bottom, notes)
    full(s, A / asset)

s = new_slide("Real micrographs: results", "Circumferential connectivity and closure rise, δ½ falls 88 → 24 µm; through-wall connectivity stays low",
              "The v1 values separate effects the single prototype number merged. Area fraction also rises (5 to 14 %), so this series alone cannot separate amount from connectivity; the benchmark does. At 3.2 um/px the owner's auto rule floors delta_max to one pixel (smallest accepted hydride about 10 um long).")
picture(s, A / "real_results.png", 0.25, BODY_TOP, W - 0.5, 3.35)
sp = real["specimens"]
rows = [["Specimen", "Af", "Prototype (report / re-run)", "HCI R", "HCI C", "HCI iso", "δ½ R (µm)", "κ"]]
for key, label in (("25ppm", "25 ppm"), ("70ppm", "70 ppm"), ("200ppm", "200 ppm")):
    f = sp[key]["report_mask"]["fixed_10um"]
    rows.append([
        label,
        f"{100 * sp[key]['report_mask_area_fraction']:.1f} %",
        f"{real['prototype_reported_hci'][key]:.2f} / {sp[key]['prototype_on_report_mask']['workbook_hci']:.2f}",
        f"{f['HCI']['radial']:.3f}", f"{f['HCI']['circumferential']:.3f}", f"{f['HCI']['isotropic']:.3f}",
        f"{f['delta_half']['radial']:.0f}", f"{f['topology']['network_closure']:.2f}",
    ])
table(s, 0.4, BODY_TOP + 3.45, W - 0.8, rows, [1.6, 1.1, 3.2, 1.2, 1.2, 1.3, 1.45, 1.48], row_h=0.5, bold_col0=True)

# Implementation
s = new_slide("Implementation: enabled by default in the web UI", "One library call behind the web UI, CLI and Python; results in JSON, PDF and XLSX",
              "Web: section 5 'Connectivity (HCI)' is ticked by default; delta_max mode (auto smallest / auto median / fixed), optional pixel size and radial reference length. Results panel shows HCI radial, circumferential and isotropic with delta_half, delta_max, counts and quality flags; views for clusters, curve, topology and path; a PDF page and an XLSX Connectivity sheet. CLI: microseg_cli.py hci. Python: analyze_continuity().")
picture(s, A / "ui_hci_controls.png", 0.3, BODY_TOP + 0.05, 3.9, 3.2)
picture(s, A / "ui_hci_panel.png", 4.35, BODY_TOP + 0.05, 8.7, 1.85)
picture(s, A / "ui_hci_curve_view.png", 4.35, BODY_TOP + 2.0, 8.7, BODY_BOTTOM - BODY_TOP - 2.05)
textbox(s, 0.3, BODY_TOP + 3.35, 3.95, 2.4, [("Also", ["CLI: microseg_cli.py hci", "Python: analyze_continuity()", "PDF page + XLSX sheet"], TEAL)])

s = new_slide("Documentation: theory and algorithm with typeset mathematics", "The same equations appear in the web Help (KaTeX) and the Sphinx docs (MathJax), with proofs",
              "Web Help section 'Hydride Connectivity Index (HCI)': definitions, connectivity function, delta_max rule with the comparison warning, companions, properties, algorithm, settings and reporting rules. Sphinx: docs/hci_theory.md (propositions and proofs, closure identity, algorithm and complexity), docs/hci_user_guide.md, docs/hci_specification.md, docs/hci_synthetic_benchmark.md.")
picture(s, A / "ui_help_math.png", 0.3, BODY_TOP, 8.2, BODY_BOTTOM - BODY_TOP)
textbox(s, 8.7, BODY_TOP + 0.2, 4.35, 5.5, [
    ("Web Help", ["Rendered offline with KaTeX", "Concept + flow SVGs"], TEAL),
    ("Sphinx", ["hci_theory: proofs", "hci_user_guide", "hci_specification", "hci_synthetic_benchmark"], NAVY),
    ("Tests", ["40 HCI tests", "Axioms, topology, parity"], MAROON),
])

# Caveats
s = new_slide("Caveats and limitations", "HCI is experimental: validated on synthetic axioms, awaiting field data and property correlation",
              "Each caveat is also flagged in the result JSON where it can be detected automatically: uncalibrated lengths, largest cluster touching the border, extent capped at the reference length, hydride thickness below 3 px, clean-up thresholds raised to the resolution floor, topology identity residual non-zero. The web panel shows the flags (right).")
textbox(s, 0.4, BODY_TOP + 0.1, 5.6, 5.7, [
    ("Comparability", ["Auto δ_max differs per image: fix δ_max to compare", "Same pixel size and reference length"], MAROON),
    ("Physics", ["2D section: out-of-plane links unseen", "Depends on Af across morphologies", "No mechanical correlation yet"], TEAL),
    ("Image quality", ["Border clusters truncated (lower bound)", "Thin masks (< 3 px): topology unreliable"], OCHRE),
    ("Performance", ["2048² mask ≈ 8 s on one core"], NAVY),
], size=18)
picture(s, A / "ui_hci_panel.png", 6.2, BODY_TOP + 0.1, W - 6.45, 2.0)
picture(s, A / "real200_stage3.png", 6.2, BODY_TOP + 2.3, W - 6.45, BODY_BOTTOM - BODY_TOP - 2.35)

# Summary
s = new_slide("Summary and next steps", "HCI v1 is specified, validated on synthetic data, implemented, and ready for field-data validation",
              "Next: field-data validation by the owner's team (repeat fields, bootstrap CIs, predicted versus corrected masks, expert ranking, correlation with ring compression, DBTT or DHC), the desktop (Qt) card, and a faster path solver. Manuscript outline in docs/hci_implementation_plan.md.")
textbox(s, 0.4, BODY_TOP + 0.1, 4.6, 5.7, [
    ("Delivered", ["Specification, theory, proofs", "203-mask constant-Af benchmark", "Library, CLI, web (on by default)", "40 tests; exact topology", "Report images analysed"], TEAL),
    ("Next", ["Field-data validation", "Fixed-δ_max comparison protocol", "Desktop card; faster path solver", "Manuscript"], MAROON),
], size=18)
picture(s, A / "problem_same_af.png", 5.1, BODY_TOP + 0.1, W - 5.35, BODY_BOTTOM - BODY_TOP - 0.2)

prs.save(OUT)
print(f"deck: {OUT} ({len(prs.slides)} slides)")
